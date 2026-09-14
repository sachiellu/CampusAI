from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
import os

out_dir = r"C:\C_projects\CampusAI-Agent\docs"
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "CampusAI-MCP-Deck.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x0F, 0x1C, 0x24)
ACCENT = RGBColor(0x2E, 0xC4, 0xB6)
TEXT = RGBColor(0xED, 0xF4, 0xF6)
MUTED = RGBColor(0x9D, 0xB4, 0xBD)


def set_run_font(run, size=20, bold=False, color=TEXT, name="Microsoft JhengHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{name}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", name)


def add_bg(slide):
    shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.15), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_title(slide, text, top=0.35, size=32):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.2), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=True, color=ACCENT)


def add_body(slide, lines, top=1.3, size=20):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.2), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = line
        set_run_font(run, size=size, bold=False, color=TEXT)


def add_footer(slide, page, total=9):
    box = slide.shapes.add_textbox(Inches(11.2), Inches(7.05), Inches(1.8), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{page} / {total}"
    set_run_font(run, size=12, color=MUTED)


def new_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide)
    return slide


slides_data = [
    {
        "title": "校園學生服務個人化 AI 助理",
        "lines": [
            "MCP Client／Server 架構｜Server 僅 RAG 檢索｜Client 負責對話管理",
            "",
            "技術棧：Vue ＋ ASP.NET Core ＋ JWT ＋ SQL Server ＋ Gemini ＋ MCP",
            "作品定位：可 Demo 的真實系統（非示意投影片）",
            "",
            "作者：（請填寫你的名字）",
        ],
    },
    {
        "title": "題目規格對照",
        "lines": [
            "規格：完成 MCP Client／Server 架構 → 獨立 MCP RAG Server ＋ API 作為 MCP Client",
            "規格：Server 端只有 RAG 檢索 → 唯一 Tool＝rag_retrieve",
            "規格：Client 端負責對話管理 → Vue 多對話側欄 ＋ JWT ＋ SQL 持久化",
            "規格：完成後以簡報說明設計及與規格 → 本簡報",
            "熟悉語言實作 → C#／TypeScript（Vue）",
        ],
    },
    {
        "title": "系統架構（一句話）",
        "lines": [
            "Client 管聊天；Server 不聊天，只檢索。",
            "",
            "使用者 → Vue（對話／多對話側欄）",
            "→ CampusAI API（Orchestrator、JWT、SQL）",
            "→ 校園問答時呼叫 MCP Client",
            "→ MCP RAG Server（rag_retrieve）",
            "→ 知識庫 knowledge／campus-docs.json",
            "",
            "選課／獎學金／活動：校務 Tools｜問答：MCP RAG",
        ],
    },
    {
        "title": "MCP Server 設計",
        "lines": [
            "專案：mcp-rag-server（埠 5099）",
            "協定：MCP over HTTP",
            "唯一 Tool：rag_retrieve(query, topK)",
            "輸入：自然語言問題｜輸出：Top-K 校園文件（標題、類別、內容、分數）",
            "知識來源：獎學金總則、選課注意、活動辦法等",
            "",
            "明確不做：對話狀態、登入、LLM 回覆",
        ],
    },
    {
        "title": "MCP Client／對話管理",
        "lines": [
            "登入（JWT）、角色（學生／教師／行政）",
            "主畫面選功能；左側：新對話／歷史／刪除／回主畫面",
            "對話持久化：SQL Server（SSMS 可查 ChatConversations）",
            "需要知識時才呼叫 MCP rag_retrieve",
            "檢索結果再交給 Gemini 潤成自然語言",
            "工具呼叫紀錄可在 UI 展開驗證 MCP 軌跡",
        ],
    },
    {
        "title": "一輪請求怎麼走",
        "lines": [
            "1. 使用者在「校園問答」輸入問題",
            "2. Client 送出（帶 conversationId）",
            "3. API 判斷 intent＝校園問答",
            "4. MCP Client 呼叫 rag_retrieve",
            "5. 取回相關文件",
            "6. LLM 依文件生成回覆",
            "7. 存進該對話，畫面顯示人話＋工具紀錄",
        ],
    },
    {
        "title": "Demo 重點（請替換成截圖）",
        "lines": [
            "建議放三張截圖：",
            "① 主畫面 ＋ 左側對話側欄",
            "② 問「獎學金申請辦法是什麼？」的回覆",
            "③ 工具呼叫紀錄出現 MCP Client／rag_retrieve",
            "",
            "口頭一句：Server 只檢索，聊天與歷史都在 Client。",
            "（此頁可刪文案、改貼圖）",
        ],
    },
    {
        "title": "與規格的對齊結論",
        "lines": [
            "MCP Client／Server → 已達成",
            "Server 僅 RAG（單一 tool）→ 已達成",
            "Client 對話管理（多對話＋JWT＋SQL）→ 已達成",
            "簡報說明設計與規格 → 本簡報",
            "",
            "延伸：可把關鍵字 RAG 換成向量 Embedding／向量庫",
            "介面仍可維持同一個 MCP tool，Client 不用大改",
        ],
    },
    {
        "title": "總結",
        "lines": [
            "用 MCP 把「檢索」與「對話」解耦",
            "Server 專注 RAG，Client 專注產品與狀態",
            "真實可 Demo：登入、多對話、MCP 工具軌跡、SQL 持久化",
            "",
            "感謝聆聽｜歡迎提問",
        ],
    },
]

for i, s in enumerate(slides_data, 1):
    slide = new_slide()
    add_title(slide, s["title"], size=30 if i > 1 else 34)
    add_body(slide, s["lines"], top=1.25 if i > 1 else 1.4, size=18 if i != 1 else 20)
    add_footer(slide, i)

prs.save(out_path)
print(out_path)
print("slides", len(prs.slides))
