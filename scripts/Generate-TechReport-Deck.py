"""Campus Personalized AI Assistant — Technical Report deck (dark theme, current stack)."""
from __future__ import annotations

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

ROOT = r"C:\C_projects\CampusAI-Agent"
OUT = os.path.join(ROOT, "docs", "CampusAI-TechReport.pptx")

SW, SH = 13.333, 7.5
ML = 0.55

BG = RGBColor(0x0B, 0x12, 0x1A)
PANEL = RGBColor(0x12, 0x1C, 0x28)
PANEL2 = RGBColor(0x15, 0x22, 0x30)
CYAN = RGBColor(0x2E, 0xE6, 0xD6)
CYAN_DIM = RGBColor(0x1A, 0xA8, 0x9C)
GREEN = RGBColor(0x3D, 0xD6, 0x8C)
ORANGE = RGBColor(0xF0, 0xA0, 0x4B)
RED = RGBColor(0xFF, 0x6B, 0x6B)
WHITE = RGBColor(0xF2, 0xF5, 0xF7)
MUTED = RGBColor(0x8A, 0x9A, 0xA8)
LINE = RGBColor(0x2A, 0x3D, 0x52)
GLOW = RGBColor(0x1C, 0x3A, 0x48)


def font(run, size=14, bold=False, color=WHITE):
    name = "Microsoft JhengHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{name}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", name)


def tb(slide, l, t, w, h, lines, size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT, after=4):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = line
        font(r, size=size, bold=bold, color=color)
    return box


def rect(slide, l, t, w, h, fill, line=None, width=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(width)
    return s


def bar(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    return s


def chrome(slide, page, total=9):
    tb(slide, ML, 0.22, 10.5, 0.28, ["CAMPUS PERSONALIZED AI ASSISTANT — TECHNICAL REPORT"], size=11, bold=True, color=CYAN)
    tb(slide, 11.4, 0.2, 1.5, 0.3, [f"{page:02d} SLIDE"], size=11, bold=True, color=MUTED, align=PP_ALIGN.RIGHT)


def title_block(slide, title, subtitle):
    tb(slide, ML, 0.55, 12, 0.5, [title], size=30, bold=True, color=WHITE)
    if subtitle:
        tb(slide, ML, 1.1, 12, 0.4, [subtitle], size=13, color=MUTED)


def pill(slide, l, t, w, h, text, fill=GLOW, color=CYAN):
    rect(slide, l, t, w, h, fill, CYAN, 1.25)
    tb(slide, l + 0.08, t + 0.08, w - 0.16, h - 0.1, [text], size=11, bold=True, color=color, align=PP_ALIGN.CENTER)


def flow_box(slide, l, t, w, h, title, body, tag=None, glow=False):
    border = CYAN if glow else LINE
    rect(slide, l, t, w, h, PANEL, border, 1.5 if glow else 1)
    tb(slide, l + 0.1, t + 0.1, w - 0.2, 0.35, [title], size=13, bold=True, color=CYAN if glow else WHITE)
    tb(slide, l + 0.1, t + 0.42, w - 0.2, h - 0.55, [body], size=11, color=MUTED)
    if tag:
        tb(slide, l + 0.1, t + h - 0.32, w - 0.2, 0.25, [tag], size=10, bold=True, color=GREEN)


def dashed_panel(slide, l, t, w, h):
    return rect(slide, l, t, w, h, PANEL, CYAN_DIM, 1.25)


# ---------- slides ----------

def slide_cover(prs):
    s = blank(prs)
    chrome(s, 1)
    tb(s, ML, 2.2, 12, 0.4, ["TECHNICAL REPORT"], size=14, bold=True, color=CYAN)
    tb(s, ML, 2.7, 12, 0.8, ["校園個人化助理"], size=40, bold=True, color=WHITE)
    tb(s, ML, 3.55, 12, 0.4, ["Campus Personalized AI Assistant"], size=18, color=MUTED)
    bar(s, ML, 4.15, 2.2, 0.06, CYAN)
    tb(s, ML, 4.45, 12, 0.35, ["報告人：呂紹瑜"], size=16, bold=True, color=WHITE)
    tb(s, ML, 4.9, 12, 0.35, ["React · ASP.NET Core · SQL Server · MCP RAG · Gemini"], size=13, color=MUTED)
    tb(s, ML, 6.7, 12, 0.3, ["Client 管對話與流程　｜　MCP Server 只檢索　｜　校務資料在 SQL"], size=12, color=GREEN)


def slide_spec(prs):
    s = blank(prs)
    chrome(s, 2)
    title_block(s, "題目規格對照", "嚴格遵循設計指標：MCP 端點職責分離，並對齊目前開發實作")

    cols = [
        (
            0.55,
            "MCP Server 規格與實作",
            "規格：MCP Server 只有一個功能——RAG 檢索",
            "實作：rag_retrieve(query, topK)",
            "僅開放單一端點，獨立於對話機制外；針對 campus-docs.json 做語意／關鍵字檢索。",
            GREEN,
        ),
        (
            4.55,
            "MCP Client 規格與實作",
            "規格：MCP Client 負責使用者對話管理",
            "實作：React 多對話 ＋ 後端 JWT／SQL",
            "ASP.NET Core 管理 JWT 身分；對話歷程持久化於 SQL Server；校園問答時才呼叫 MCP。",
            ORANGE,
        ),
        (
            8.55,
            "開發環境與語言",
            "本簡報：完整結構化架構解說",
            "熟悉實作語言",
            "C#（ASP.NET Core／SSMS）\nTypeScript（React ＋ Vite 介面）\nMCP RAG Server：C#（mcp-rag-server）",
            CYAN,
        ),
    ]
    for l, title, req, impl, detail, accent in cols:
        dashed_panel(s, l, 1.7, 3.8, 5.0)
        tb(s, l + 0.2, 1.9, 3.4, 0.4, [title], size=15, bold=True, color=WHITE)
        bar(s, l + 0.2, 2.4, 1.4, 0.04, accent)
        tb(s, l + 0.2, 2.65, 3.4, 0.7, [req], size=12, color=MUTED)
        tb(s, l + 0.2, 3.45, 3.4, 0.55, [impl], size=14, bold=True, color=accent)
        tb(s, l + 0.2, 4.2, 3.4, 2.2, [detail], size=12, color=WHITE)


def slide_architecture(prs):
    s = blank(prs)
    chrome(s, 3)
    title_block(s, "系統架構", "職責明確：對話與流程控制歸於 Client；知識檢索隔離於 MCP Server")

    rect(s, ML, 1.65, 12.2, 0.45, RGBColor(0x0F, 0x2A, 0x22), GREEN, 1)
    tb(s, ML + 0.2, 1.72, 11.8, 0.35, ["核心理念：Client 管聊天與調度；MCP Server 不聊天，只檢索。"], size=13, bold=True, color=GREEN)

    boxes = [
        (0.4, "使用者", "瀏覽器操作網站與對話", "Student / Teacher / Admin", False),
        (2.45, "React Frontend", "主畫面、側欄、多對話、權限入口", "TypeScript + Vite", True),
        (4.5, "CampusAI API", "JWT、對話持久化、校務 API", "ASP.NET Core", False),
        (6.55, "Personal Orchestrator", "意圖判斷、六類 Agent 調度", "Agent 層", True),
        (8.6, "MCP RAG Server", "唯一功能：rag_retrieve", "C# · MCP", False),
        (10.65, "資料來源", "SQL 校務表 ＋ campus-docs.json", "SQL / JSON", False),
    ]
    for l, title, body, tag, glow in boxes:
        flow_box(s, l, 2.4, 1.95, 2.55, title, body, tag, glow)

    for i in range(5):
        x = 0.4 + 1.95 * (i + 1) - 0.12
        bar(s, x, 3.55, 0.22, 0.04, CYAN_DIM)

    rect(s, ML, 5.3, 12.2, 1.5, PANEL, ORANGE, 1.25)
    tb(s, ML + 0.25, 5.45, 11.7, 0.35, ["關鍵分流策略"], size=14, bold=True, color=ORANGE)
    tb(
        s,
        ML + 0.25,
        5.9,
        11.7,
        0.7,
        [
            "獎學金／選課／活動等校務操作 → 直接走校務 Tools，讀寫 SQL Server。",
            "校園規章、辦法等知識問答 → 才走 MCP rag_retrieve，再由 Gemini 整理回覆。",
        ],
        size=13,
        color=WHITE,
        after=6,
    )


def slide_layers(prs):
    s = blank(prs)
    chrome(s, 4)
    title_block(s, "系統分層", "由上而下：應用 → Agent → API → 資料，對齊目前程式結構")

    layers = [
        ("應用層", "React：主畫面、校務系統、平台服務、助理功能", CYAN),
        ("Agent 層", "Personal Orchestrator ＋ 推薦／規劃／申請／溝通／提醒／分析", GREEN),
        ("API 層", "校務 API、平台服務、JWT、對話與 MCP Client", ORANGE),
        ("資料層", "SQL Server 正式校務表；校園問答另接 MCP 找文件", MUTED),
    ]
    y = 1.75
    for title, body, accent in layers:
        rect(s, ML, y, 12.2, 1.1, PANEL, accent, 1.5)
        bar(s, ML, y, 0.12, 1.1, accent)
        tb(s, ML + 0.4, y + 0.22, 3.2, 0.4, [title], size=18, bold=True, color=WHITE)
        tb(s, ML + 3.6, y + 0.3, 8.5, 0.5, [body], size=14, color=MUTED)
        y += 1.25


def slide_mcp_server(prs):
    s = blank(prs)
    chrome(s, 5)
    title_block(s, "MCP Server 設計", "極簡、無狀態的知識索取端點，單一職責")

    dashed_panel(s, 0.55, 1.7, 5.9, 5.1)
    tb(s, 0.8, 1.9, 5.4, 0.35, ["服務基本配置"], size=16, bold=True, color=WHITE)
    rows = [
        ("專案名稱", "mcp-rag-server"),
        ("通訊協定", "MCP over HTTP"),
        ("單一開放 Tool", "rag_retrieve(query, topK)"),
        ("知識來源", "knowledge/campus-docs.json"),
    ]
    yy = 2.4
    for k, v in rows:
        tb(s, 0.8, yy, 2.2, 0.3, [k], size=12, color=MUTED)
        tb(s, 2.9, yy, 3.2, 0.3, [v], size=13, bold=True, color=CYAN)
        yy += 0.45

    tb(s, 0.8, 4.4, 5.4, 0.35, ["Server 絕對不做的事"], size=15, bold=True, color=WHITE)
    for i, line in enumerate(
        [
            "不維護對話狀態",
            "不處理 JWT 與帳號登入",
            "不呼叫 LLM、不生成自然語言回覆",
        ]
    ):
        tb(s, 0.8, 4.9 + i * 0.4, 5.4, 0.35, [f"✕  {line}"], size=13, color=RED)

    dashed_panel(s, 6.7, 1.7, 6.0, 5.1)
    tb(s, 6.95, 1.9, 5.5, 0.35, ["Tool API 規格"], size=16, bold=True, color=WHITE)
    tb(s, 6.95, 2.4, 5.5, 0.3, ["輸入"], size=13, bold=True, color=CYAN)
    tb(
        s,
        6.95,
        2.75,
        5.5,
        0.9,
        ["query：自然語言問題", "topK：回傳最相關文件片段數（預設 3）"],
        size=13,
        color=WHITE,
        after=6,
    )
    tb(s, 6.95, 3.8, 5.5, 0.3, ["輸出"], size=13, bold=True, color=CYAN)
    tb(
        s,
        6.95,
        4.15,
        5.5,
        1.5,
        ["Top-K 規章片段：標題、類別、內容、分數", "僅回傳原始檢索結果，不做聊天包裝"],
        size=13,
        color=WHITE,
        after=6,
    )
    rect(s, 6.95, 5.8, 5.5, 0.7, RGBColor(0x0F, 0x2A, 0x22), GREEN, 1)
    tb(s, 7.15, 5.95, 5.1, 0.4, ["一句話：Server 只檢索，不聊天。"], size=14, bold=True, color=GREEN)


def slide_mcp_client(prs):
    s = blank(prs)
    chrome(s, 6)
    title_block(s, "MCP Client／對話管理", "統籌身分、對話持久化，並在需要時調度 MCP 與 Gemini")

    dashed_panel(s, 0.55, 1.7, 6.0, 5.1)
    tb(s, 0.8, 1.9, 5.5, 0.35, ["對話大腦與身分持久化"], size=16, bold=True, color=WHITE)
    tb(s, 0.8, 2.45, 5.5, 0.3, ["JWT 身分驗證"], size=14, bold=True, color=CYAN)
    tb(
        s,
        0.8,
        2.85,
        5.5,
        1.0,
        ["學生／教師／行政角色分流；入口與功能依身分切換。"],
        size=13,
        color=WHITE,
    )
    tb(s, 0.8, 3.9, 5.5, 0.3, ["SQL Server 對話持久化"], size=14, bold=True, color=CYAN)
    tb(
        s,
        0.8,
        4.3,
        5.5,
        1.4,
        [
            "後端寫入對話紀錄，可用 SSMS 查核。",
            "前端側欄：新對話、切換多對話、歷史清單。",
            "校務正式資料（成績、獎助、課程…）同庫管理。",
        ],
        size=13,
        color=WHITE,
        after=8,
    )

    dashed_panel(s, 6.8, 1.7, 5.9, 5.1)
    tb(s, 7.05, 1.9, 5.4, 0.35, ["知識整合與生成流程"], size=16, bold=True, color=WHITE)
    steps = [
        ("1. 意圖評估", "涉及校園規章等知識時，才呼叫 MCP rag_retrieve。"),
        ("2. 校務 Tools", "獎學金／選課／活動等直接查 SQL，不經 MCP。"),
        ("3. Gemini 整理", "把查得結果與對話脈絡組成回覆。"),
    ]
    yy = 2.5
    for title, body in steps:
        tb(s, 7.05, yy, 5.4, 0.3, [title], size=13, bold=True, color=GREEN)
        tb(s, 7.05, yy + 0.32, 5.4, 0.55, [body], size=12, color=WHITE)
        yy += 0.95
    rect(s, 7.05, 5.7, 5.4, 0.8, GLOW, CYAN, 1)
    tb(
        s,
        7.2,
        5.85,
        5.1,
        0.55,
        ["RAG 與對話解耦：檢索交給 MCP；安全與歷史交給 ASP.NET Core。"],
        size=12,
        bold=True,
        color=CYAN,
    )


def slide_request_flow(prs):
    s = blank(prs)
    chrome(s, 7)
    title_block(s, "一輪請求怎麼走", "從輸入問題到畫面回覆的完整生命週期")

    steps = [
        ("1 輸入", "使用者在網站提問"),
        ("2 送出", "帶 conversationId 與 JWT"),
        ("3 意圖", "Orchestrator 判斷路徑"),
        ("4 分流", "SQL 校務 或 MCP 檢索"),
        ("5 取回", "成績／規章片段"),
        ("6 生成", "Gemini 整理回覆"),
        ("7 落地", "SQL 存檔並渲染畫面"),
    ]
    for i, (title, body) in enumerate(steps):
        l = 0.4 + i * 1.82
        glow = i == 6
        flow_box(s, l, 1.85, 1.7, 2.3, title, body, None, glow)
        if i < 6:
            bar(s, l + 1.7, 2.9, 0.12, 0.04, CYAN_DIM)

    rect(s, ML, 4.5, 12.2, 2.2, PANEL, LINE, 1)
    tb(s, ML + 0.3, 4.7, 11.6, 0.35, ["兩條實際路徑"], size=15, bold=True, color=WHITE)
    tb(
        s,
        ML + 0.3,
        5.2,
        11.6,
        1.2,
        [
            "問獎學金／選課／活動 → API → 校務 Tools → SQL Server → Gemini 整理。",
            "問校規／辦法 → API → MCP Client → rag_retrieve → campus-docs.json → Gemini 整理。",
            "執行紀錄可看到工具呼叫軌跡（例如 Call tool: rag_retrieve）。",
        ],
        size=14,
        color=MUTED,
        after=8,
    )


def slide_demo(prs):
    s = blank(prs)
    chrome(s, 8)
    title_block(s, "現場示範", "用同一套資料，對照資料庫、畫面與執行紀錄")

    cards = [
        ("1. 資料庫", "SSMS 打開 Students，先看 GPA 與正式校務表。", "SQL Server"),
        ("2. 主畫面", "登入後走一圈：校務系統、平台服務、助理功能三塊入口。", "React"),
        ("3. 精準問答", "問獎學金條件，對照資料庫是否一致；問校規可見 MCP 軌跡。", "RAG / Tools"),
    ]
    for i, (title, body, tag) in enumerate(cards):
        l = 0.55 + i * 4.15
        dashed_panel(s, l, 1.75, 3.95, 3.4)
        tb(s, l + 0.25, 2.0, 3.45, 0.4, [title], size=16, bold=True, color=CYAN)
        tb(s, l + 0.25, 2.6, 3.45, 1.8, [body], size=14, color=WHITE, after=8)
        pill(s, l + 0.25, 4.55, 2.2, 0.35, tag)

    rect(s, ML, 5.5, 12.2, 1.3, RGBColor(0x0F, 0x2A, 0x22), GREEN, 1.25)
    tb(s, ML + 0.3, 5.7, 11.6, 0.35, ["KEY CONCEPT"], size=12, bold=True, color=GREEN)
    tb(
        s,
        ML + 0.3,
        6.15,
        11.6,
        0.45,
        ["入口依身分與功能分；校務資料在 SQL；問答找文件走 MCP；回覆由 Gemini 整理。"],
        size=15,
        bold=True,
        color=WHITE,
    )


def slide_closing(prs):
    s = blank(prs)
    chrome(s, 9)
    title_block(s, "總結", "這套系統現在的完整樣貌")

    points = [
        ("前端", "React + TypeScript + Vite，依身分切換入口"),
        ("後端", "ASP.NET Core：JWT、Orchestrator、校務 API、MCP Client"),
        ("資料", "SQL Server 正式表；校園問答用 campus-docs.json"),
        ("檢索", "MCP Server 單一工具 rag_retrieve，不聊天"),
        ("生成", "Gemini 把查得結果整理成可讀回覆"),
    ]
    y = 1.75
    for title, body in points:
        rect(s, ML, y, 12.2, 0.8, PANEL, LINE, 1)
        tb(s, ML + 0.35, y + 0.2, 2.2, 0.4, [title], size=16, bold=True, color=CYAN)
        tb(s, ML + 2.8, y + 0.22, 9.5, 0.4, [body], size=15, color=WHITE)
        y += 0.95

    tb(s, ML, 6.7, 12, 0.35, ["報告人：呂紹瑜　｜　感謝聆聽"], size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    slide_cover(prs)
    slide_spec(prs)
    slide_architecture(prs)
    slide_layers(prs)
    slide_mcp_server(prs)
    slide_mcp_client(prs)
    slide_request_flow(prs)
    slide_demo(prs)
    slide_closing(prs)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
