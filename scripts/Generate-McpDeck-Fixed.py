"""Campus AI 簡報：明亮正式＋圖表排版；講用途與架構，不講「做／不做」。"""
from __future__ import annotations

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from PIL import Image, ImageDraw

ROOT = r"C:\C_projects\CampusAI-Agent"
ASSETS = os.path.join(ROOT, "docs", "assets")
OUT = os.path.join(ROOT, "docs", "CampusAI-MCP-Deck-v3.pptx")
BG = os.path.join(ASSETS, "slide-bg-campus-plain.png")

SW, SH = 13.333, 7.5
ML, MT = 0.55, 0.32
SAFE_W = 12.2
SAFE_RIGHT = ML + SAFE_W

CREAM = (247, 245, 240)
INK = RGBColor(0x1C, 0x24, 0x28)
MUTED = RGBColor(0x5A, 0x65, 0x6B)
TEAL = RGBColor(0x1F, 0x8A, 0x80)
TEAL_SOFT = RGBColor(0xE6, 0xF4, 0xF2)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD4, 0xD8, 0xD6)
ORANGE = RGBColor(0xC4, 0x7A, 0x3A)
GREEN = RGBColor(0x2F, 0x7A, 0x4F)


def make_bg(path: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    img = Image.new("RGB", (1920, 1080), CREAM)
    d = ImageDraw.Draw(img)
    for y in range(48, 1080, 40):
        for x in range(48, 1920, 40):
            d.ellipse((x, y, x + 2, y + 2), fill=(205, 202, 194))
    img.save(path, "PNG")


def font(run, size=16, bold=False, color=INK):
    name = "Microsoft JhengHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{name}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", name)


def tb(slide, l, t, w, h, lines, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, after=5):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = line
        font(r, size=size, bold=bold, color=color)
    return box


def card(slide, l, t, w, h, fill=PAPER):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = LINE
    s.line.width = Pt(1)
    return s


def header(slide, title, subtitle, page, total=9):
    tb(slide, ML, MT, 9.5, 0.28, ["Campus AI"], size=12, bold=True, color=TEAL)
    tb(slide, ML, MT + 0.28, 10.5, 0.48, [title], size=26, bold=True, color=INK)
    if subtitle:
        tb(slide, ML, MT + 0.78, 10.5, 0.32, [subtitle], size=13, color=MUTED)
    tb(slide, 11.7, MT + 0.1, 1.0, 0.3, [f"{page}/{total}"], size=12, color=MUTED, align=PP_ALIGN.RIGHT)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(BG, 0, 0, width=prs.slide_width, height=prs.slide_height)
    return s


def main():
    make_bg(BG)
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    N = 9

    # 1 封面
    s = blank(prs)
    tb(s, ML, 1.9, 11, 0.35, ["Campus AI"], size=15, bold=True, color=TEAL)
    tb(s, ML, 2.4, 11.5, 0.9, ["校園學生服務個人化 AI 助理"], size=34, bold=True, color=INK)
    tb(
        s,
        ML,
        3.5,
        11.5,
        0.9,
        [
            "提供學生、教師、行政使用的校園助理：處理獎學金、選課、活動與校務相關問題。",
            "系統先取得校務資料與文件，再由 Gemini 整理成回覆。",
        ],
        size=16,
        color=MUTED,
        after=8,
    )
    techs = ["React", "ASP.NET Core", "SQL Server", "JWT", "Gemini", "MCP"]
    x = ML
    for t in techs:
        w = 1.55 if len(t) < 10 else 1.9
        card(s, x, 5.0, w, 0.42, fill=TEAL_SOFT)
        tb(s, x, 5.08, w, 0.3, [t], size=12, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        x += w + 0.12
    tb(s, ML, 6.4, 8, 0.35, ["簡報人：（請填姓名）"], size=14, color=INK)

    # 2 這套系統在做什麼
    s = blank(prs)
    header(s, "系統用途", "誰用、解決什麼問題", 2, N)
    cols = [
        ("學生", "獎學金資格、選課建議、校園活動", GREEN),
        ("教師", "導生關懷、成績與畢業門檻提醒", TEAL),
        ("行政", "校務統計摘要、決策支援資訊", ORANGE),
    ]
    cw = (SAFE_W - 0.4) / 3
    for i, (title, body, color) in enumerate(cols):
        left = ML + i * (cw + 0.2)
        card(s, left, 1.55, cw, 3.4)
        tb(s, left + 0.2, 1.85, cw - 0.4, 0.45, [title], size=20, bold=True, color=color)
        tb(s, left + 0.2, 2.55, cw - 0.4, 2.0, [body], size=15, color=INK)
    card(s, ML, 5.2, SAFE_W, 1.55, fill=TEAL_SOFT)
    tb(
        s,
        ML + 0.3,
        5.5,
        SAFE_W - 0.6,
        1.0,
        [
            "同一個入口依身分顯示不同功能。",
            "使用者用對話提出需求，系統回傳可執行的資訊與建議。",
        ],
        size=16,
        color=INK,
        after=6,
    )

    # 3 架構圖
    s = blank(prs)
    header(s, "系統架構", "應用層 → Agent → API → 資料", 3, N)
    layers = [
        ("應用層", "React：主畫面、校務系統、平台服務、AI 功能"),
        ("Agent 層", "Orchestrator 調度推薦／規劃／申請／溝通／提醒／分析"),
        ("API 服務層", "校務 API（學生、課程、活動…）與平台服務"),
        ("資料層", "SQL Server 校務資料表；校園問答另接 MCP RAG"),
    ]
    y = 1.5
    for title, body in layers:
        card(s, ML, y, SAFE_W, 1.15)
        tb(s, ML + 0.25, y + 0.2, 2.3, 0.4, [title], size=15, bold=True, color=TEAL)
        tb(s, ML + 2.7, y + 0.25, SAFE_W - 3.1, 0.7, [body], size=15, color=INK)
        y += 1.28

    # 4 產品入口為什麼這樣切
    s = blank(prs)
    header(s, "產品入口設計", "主畫面下的三個區塊各自負責什麼", 4, N)
    items = [
        ("主畫面", "總入口，導向三大區塊"),
        ("校務系統", "學生、課程、活動、申請、行事曆等校務功能入口"),
        ("平台服務", "身分驗證、權限、Tool Calling 等平台能力"),
        ("AI 功能", "問答、獎學金、選課、活動等助理對話"),
    ]
    cw = (SAFE_W - 0.3) / 2
    for i, (title, body) in enumerate(items):
        left = ML + (i % 2) * (cw + 0.3)
        top = 1.5 + (i // 2) * 2.5
        card(s, left, top, cw, 2.25)
        tb(s, left + 0.25, top + 0.35, cw - 0.5, 0.45, [title], size=18, bold=True, color=TEAL)
        tb(s, left + 0.25, top + 1.0, cw - 0.5, 0.9, [body], size=15, color=INK)

    # 5 技術棧與資料
    s = blank(prs)
    header(s, "技術棧與資料", "前端、後端、資料庫怎麼分工", 5, N)
    card(s, ML, 1.5, 5.8, 5.3)
    tb(s, ML + 0.25, 1.75, 5.3, 0.4, ["技術棧"], size=16, bold=True, color=TEAL)
    tb(
        s,
        ML + 0.25,
        2.35,
        5.3,
        4.0,
        [
            "前端：React + TypeScript + Vite",
            "後端：C# ASP.NET Core",
            "驗證：JWT（學生／教師／行政）",
            "資料庫：SQL Server（CampusAI）",
            "語言模型：Gemini",
            "文件檢索：MCP RAG Server",
            "通訊：REST、JSON camelCase",
        ],
        size=15,
        color=INK,
        after=10,
    )
    card(s, ML + 6.05, 1.5, SAFE_W - 6.05, 5.3)
    tb(s, ML + 6.3, 1.75, SAFE_W - 6.5, 0.4, ["SQL 主要資料表"], size=16, bold=True, color=TEAL)
    tb(
        s,
        ML + 6.3,
        2.35,
        SAFE_W - 6.5,
        4.0,
        [
            "Users｜登入帳號",
            "Students｜GPA、系級",
            "Courses／Scholarships／Activities",
            "ChatConversations／Messages",
            "",
            "助理回覆的資格與成績，",
            "來源是資料表，可在 SSMS 查看。",
        ],
        size=15,
        color=INK,
        after=8,
    )

    # 6 MCP
    s = blank(prs)
    header(s, "MCP Client／Server", "校園問答的文件檢索怎麼接", 6, N)
    boxes = [
        ("使用者", "React 對話"),
        ("CampusAI API", "JWT／Agent／Gemini"),
        ("MCP Client", "後端內呼叫"),
        ("MCP RAG Server", "rag_retrieve"),
        ("知識庫", "campus-docs"),
    ]
    bw = (SAFE_W - 0.4) / 5
    for i, (a, b) in enumerate(boxes):
        left = ML + i * (bw + 0.1)
        card(s, left, 1.6, bw, 2.6)
        tb(s, left + 0.1, 1.9, bw - 0.2, 0.8, [a], size=13, bold=True, color=TEAL)
        tb(s, left + 0.1, 2.9, bw - 0.2, 1.0, [b], size=13, color=INK)
    card(s, ML, 4.55, SAFE_W, 2.2, fill=TEAL_SOFT)
    tb(
        s,
        ML + 0.3,
        4.9,
        SAFE_W - 0.6,
        1.5,
        [
            "獎學金、選課、活動：走校務 API 與 SQL。",
            "校園問答：走 MCP 檢索規章文件。",
            "對話狀態、登入與歷史由 API 與資料庫管理。",
        ],
        size=16,
        color=INK,
        after=8,
    )

    # 7 一輪流程
    s = blank(prs)
    header(s, "一次提問的處理流程", "從輸入到回覆", 7, N)
    steps = [
        ("1", "前端送出問題（含登入權杖與對話編號）"),
        ("2", "Orchestrator 判斷功能與意圖"),
        ("3", "查 SQL 校務資料，或呼叫 MCP 檢索文件"),
        ("4", "Gemini 依查詢結果整理回覆"),
        ("5", "寫入對話紀錄，畫面顯示結果與執行紀錄"),
    ]
    y = 1.5
    for num, line in steps:
        card(s, ML, y, SAFE_W, 0.85)
        tb(s, ML + 0.25, y + 0.22, 0.5, 0.4, [num], size=18, bold=True, color=TEAL)
        tb(s, ML + 0.9, y + 0.25, SAFE_W - 1.3, 0.45, [line], size=16, color=INK)
        y += 0.98

    # 8 Demo
    s = blank(prs)
    header(s, "Demo 重點", "建議展示順序", 8, N)
    panels = [
        ("介面", "登入 → 主畫面 → 校務／平台／AI"),
        ("資料", "SSMS 開啟 Students，對照 GPA"),
        ("助理", "獎學金或校園問答，查看回覆與執行紀錄"),
    ]
    pw = (SAFE_W - 0.4) / 3
    for i, (t, b) in enumerate(panels):
        left = ML + i * (pw + 0.2)
        card(s, left, 1.55, pw, 4.5)
        tb(s, left + 0.2, 1.9, pw - 0.4, 0.5, [t], size=18, bold=True, color=TEAL)
        tb(s, left + 0.2, 2.7, pw - 0.4, 2.8, [b, "", "（可放截圖）"], size=15, color=INK)

    # 9 Q&A — 獨立專區
    s = blank(prs)
    card(s, 2.2, 2.0, 8.9, 3.4, fill=PAPER)
    tb(s, 2.2, 2.5, 8.9, 0.7, ["Q & A"], size=40, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    tb(s, 2.2, 3.5, 8.9, 0.9, ["歡迎提問", "Campus AI｜校園學生服務個人化 AI 助理"], size=16, color=MUTED, align=PP_ALIGN.CENTER, after=8)

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
