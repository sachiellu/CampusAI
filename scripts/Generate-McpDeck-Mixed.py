"""依使用者指定的白話風格重寫；混排：純文字／架構圖交替。v3-saved 不動。"""
from __future__ import annotations

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from PIL import Image, ImageDraw

ROOT = r"C:\C_projects\CampusAI-Agent"
ASSETS = os.path.join(ROOT, "docs", "assets")
OUT = os.path.join(ROOT, "docs", "CampusAI-MCP-Deck-v5.pptx")
BG = os.path.join(ASSETS, "slide-bg-campus-plain.png")

SW, SH = 13.333, 7.5
ML, MT = 0.7, 0.35
SAFE_W = 11.9

CREAM = (247, 245, 240)
INK = RGBColor(0x1C, 0x24, 0x28)
MUTED = RGBColor(0x5A, 0x65, 0x6B)
TEAL = RGBColor(0x1F, 0x8A, 0x80)
TEAL_SOFT = RGBColor(0xE6, 0xF4, 0xF2)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD4, 0xD8, 0xD6)
ORANGE = RGBColor(0xC4, 0x7A, 0x3A)
GREEN = RGBColor(0x2F, 0x7A, 0x4F)


def make_bg(path: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    img = Image.new("RGB", (1920, 1080), CREAM)
    d = ImageDraw.Draw(img)
    for y in range(48, 1080, 40):
        for x in range(48, 1920, 40):
            d.ellipse((x, y, x + 2, y + 2), fill=(205, 202, 194))
    img.save(path, "PNG")


def font(run, size=16, bold=False, color=INK):
    name = "Microsoft JhengHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(
            f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{name}"/>'
        )
        rPr.append(ea)
    else:
        ea.set("typeface", name)


def tb(slide, l, t, w, h, lines, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, after=6):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = line
        font(r, size=size, bold=bold, color=color)
    return box


def card(slide, l, t, w, h, fill=PAPER):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = LINE
    s.line.width = Pt(1)
    return s


def rule(slide, y):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(ML), Inches(y), Inches(SAFE_W), Inches(0.015))
    sh.fill.solid()
    sh.fill.fore_color.rgb = LINE
    sh.line.fill.background()


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(BG, 0, 0, width=prs.slide_width, height=prs.slide_height)
    return s


def head_text(slide, title, page, total=9):
    tb(slide, ML, MT, 10, 0.3, ["Campus AI"], size=12, bold=True, color=TEAL)
    tb(slide, ML, MT + 0.32, 10.5, 0.5, [title], size=28, bold=True, color=INK)
    tb(slide, 11.5, MT + 0.1, 1.2, 0.3, [f"{page}/{total}"], size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    rule(slide, 1.35)


def head_diagram(slide, title, subtitle, page, total=9):
    tb(slide, ML, MT, 10, 0.28, ["Campus AI"], size=12, bold=True, color=TEAL)
    tb(slide, ML, MT + 0.28, 10.5, 0.45, [title], size=26, bold=True, color=INK)
    if subtitle:
        tb(slide, ML, MT + 0.75, 10.5, 0.3, [subtitle], size=13, color=MUTED)
    tb(slide, 11.5, MT + 0.1, 1.2, 0.3, [f"{page}/{total}"], size=12, color=MUTED, align=PP_ALIGN.RIGHT)


def main():
    make_bg(BG)
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    N = 9

    # 1 封面
    s = blank(prs)
    tb(s, ML, 1.95, 11, 0.35, ["Campus AI"], size=15, bold=True, color=TEAL)
    tb(s, ML, 2.45, 11.5, 0.9, ["校園學生服務個人化 AI 助理"], size=34, bold=True, color=INK)
    tb(
        s,
        ML,
        3.55,
        11.5,
        1.0,
        [
            "依身分切換入口，用對話處理獎學金、選課、活動與校務問題。",
            "正式資料查 SQL Server；校園問答檢索規章；Gemini 整理成回覆。",
        ],
        size=16,
        color=MUTED,
        after=8,
    )
    x = ML
    for t in ["React", "ASP.NET Core", "SQL Server", "JWT", "Gemini", "MCP"]:
        w = 1.55 if len(t) < 10 else 1.95
        card(s, x, 5.15, w, 0.4, fill=TEAL_SOFT)
        tb(s, x, 5.22, w, 0.28, [t], size=11, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
        x += w + 0.1
    tb(s, ML, 6.4, 8, 0.35, ["簡報人：（請填姓名）"], size=14, color=INK)

    # 2 純文字 — 對齊使用者範例語氣
    s = blank(prs)
    head_text(s, "這套系統在做什麼？", 2, N)
    tb(
        s,
        ML,
        1.7,
        12,
        5.3,
        [
            "這是一個個人化 AI 助理。",
            "依登入身分，切到不同入口：",
            "",
            "學生：獎學金、選課、活動",
            "教師：導生關懷、門檻提醒",
            "行政：校務研究、決策支援",
            "",
            "依不同身分切換不同入口。",
            "需求用對話提出，系統查完資料後回覆。",
            "",
            "資格、成績等正式資料來自 SQL Server；",
            "一般校園問答則靠檢索規章文件。",
        ],
        size=18,
        color=INK,
        after=8,
    )

    # 3 架構圖
    s = blank(prs)
    head_diagram(s, "系統怎麼分層？", "從畫面到資料", 3, N)
    layers = [
        ("應用層", "你看到的網站：主畫面、校務系統、平台服務、AI 功能"),
        ("Agent 層", "負責調度：推薦、規劃、申請、溝通、提醒、分析"),
        ("API 層", "對外提供校務與平台服務的介面"),
        ("資料層", "SQL Server 存校務資料；問答另接 MCP 找文件"),
    ]
    y = 1.45
    for title, body in layers:
        card(s, ML, y, SAFE_W, 1.15)
        tb(s, ML + 0.25, y + 0.2, 2.2, 0.4, [title], size=15, bold=True, color=TEAL)
        tb(s, ML + 2.6, y + 0.25, SAFE_W - 3.0, 0.7, [body], size=15, color=INK)
        y += 1.28

    # 4 三欄身分
    s = blank(prs)
    head_diagram(s, "三種身分各自看什麼？", None, 4, N)
    cols = [
        ("學生", "獎學金\n選課\n校園活動", GREEN),
        ("教師", "導生關懷\n門檻提醒", TEAL),
        ("行政", "校務研究\n決策支援", ORANGE),
    ]
    cw = (SAFE_W - 0.4) / 3
    for i, (title, body, color) in enumerate(cols):
        left = ML + i * (cw + 0.2)
        card(s, left, 1.55, cw, 4.0)
        tb(s, left + 0.2, 1.95, cw - 0.4, 0.5, [title], size=22, bold=True, color=color)
        tb(s, left + 0.2, 2.8, cw - 0.4, 2.2, body.split("\n"), size=18, color=INK, after=10)
    tb(s, ML, 5.9, SAFE_W, 0.8, ["同一個網站，登入後入口不同。"], size=16, color=MUTED)

    # 5 純文字 — 入口為什麼拆
    s = blank(prs)
    head_text(s, "主畫面為什麼拆成三塊？", 5, N)
    tb(
        s,
        ML,
        1.7,
        12,
        5.3,
        [
            "主畫面是總入口。",
            "",
            "校務系統：學生資料、課程、活動、申請、行事曆。",
            "之後若要接真實校務，入口已經留好。",
            "",
            "平台服務：身分驗證、權限、Tool Calling。",
            "處理的是「誰能用、能用到哪」。",
            "",
            "AI 功能：真正能對話的助理。",
            "問答、獎學金、選課、活動都從這裡進。",
            "",
            "拆開是為了分清楚：校務入口、平台能力、助理功能，",
            "不要全部擠在同一個聊天框。",
        ],
        size=17,
        color=INK,
        after=7,
    )

    # 6 技術＋資料表
    s = blank(prs)
    head_diagram(s, "用了哪些技術？資料放哪？", None, 6, N)
    tb(
        s,
        ML,
        1.45,
        5.8,
        5.5,
        [
            "前端：React",
            "後端：ASP.NET Core（C#）",
            "登入：JWT",
            "資料庫：SQL Server",
            "回覆整理：Gemini",
            "文件檢索：MCP RAG Server",
            "",
            "例如林小芸的 GPA，",
            "存在 Students 資料表，",
            "不是寫死在網頁上。",
        ],
        size=17,
        color=INK,
        after=9,
    )
    card(s, ML + 6.2, 1.45, SAFE_W - 6.2, 5.3)
    tb(s, ML + 6.45, 1.75, SAFE_W - 6.7, 0.4, ["資料庫裡有什麼"], size=15, bold=True, color=TEAL)
    tb(
        s,
        ML + 6.45,
        2.4,
        SAFE_W - 6.7,
        4.0,
        [
            "Users　帳號",
            "Students　學生與 GPA",
            "Courses　課程",
            "Scholarships　獎學金",
            "Activities　活動",
            "對話紀錄表",
        ],
        size=16,
        color=INK,
        after=12,
    )

    # 7 MCP 圖＋白話
    s = blank(prs)
    head_diagram(s, "問問題時資料怎麼來？", None, 7, N)
    boxes = [
        ("你", "在網站提問"),
        ("後端", "判斷要查什麼"),
        ("SQL", "成績、獎助…"),
        ("MCP", "找規章文件"),
        ("Gemini", "整理成回覆"),
    ]
    bw = (SAFE_W - 0.4) / 5
    for i, (a, b) in enumerate(boxes):
        left = ML + i * (bw + 0.1)
        card(s, left, 1.5, bw, 2.35)
        tb(s, left + 0.08, 1.75, bw - 0.16, 0.55, [a], size=14, bold=True, color=TEAL)
        tb(s, left + 0.08, 2.5, bw - 0.16, 1.0, [b], size=13, color=INK)
    tb(
        s,
        ML,
        4.2,
        SAFE_W,
        2.6,
        [
            "獎學金、選課、活動：查 SQL。",
            "校園問答：走 MCP 找文件。",
            "",
            "對話紀錄存在資料庫，下次還找得到。",
        ],
        size=17,
        color=INK,
        after=10,
    )

    # 8 Demo 純文字
    s = blank(prs)
    head_text(s, "Demo 可以怎麼講？", 8, N)
    tb(
        s,
        ML,
        1.7,
        12,
        5.3,
        [
            "先打開 SSMS，給大家看 Students 的 GPA。",
            "",
            "再登入網站，走一圈主畫面的三塊入口。",
            "",
            "進 AI，問獎學金，對一下是不是跟資料庫條件對得上。",
            "",
            "若問校規，打開執行紀錄，指出有打到 MCP。",
            "",
            "收尾：入口照身分與功能分；資料在 SQL；問答找文件；回覆由 Gemini 整理。",
        ],
        size=17,
        color=INK,
        after=9,
    )

    # 9 Q&A
    s = blank(prs)
    card(s, 2.3, 2.1, 8.7, 3.2)
    tb(s, 2.3, 2.6, 8.7, 0.8, ["Q & A"], size=42, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    tb(s, 2.3, 3.6, 8.7, 0.8, ["Campus AI｜校園學生服務個人化 AI 助理"], size=15, color=MUTED, align=PP_ALIGN.CENTER)

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
